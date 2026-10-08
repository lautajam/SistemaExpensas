# -*- coding: utf-8 -*-
"""
pagos.py
--------
Planillas de pagos: un Excel por edificio en datos/pagos_edificios/, llamado
"<Nombre_Edificio>_pagos.xlsx". La app lo crea y mantiene ordenado sola
(una fila por unidad); las personas solo cargan los montos de cada mes.

Diseño de la hoja:
    fila 2:  nombre del edificio
    fila 3:  Tipo unidad | Total a pagar | Monto deuda | Monto pagado | Tipo pago
    fila 4+: una fila por unidad (deptos, locales, cocheras, bauleras), en
             orden de piso y unidad, con cada depto seguido de sus cocheras
             y bauleras. Sin unidades, la fila 4 tiene igual la fórmula de
             "Tipo pago" (para poder copiarla)
    columna H (oculta): id interno de la unidad. Es lo que vincula cada fila
    con su unidad, aunque cambie el piso/letra o se reordenen las filas.

"Tipo pago" es una fórmula de Excel. Para cambiar qué texto se imprime en
"Gastos de" por cada estado, editar TEXTO_POR_ESTADO.

Modo "celdas" (planilla propia con su orden): si alguna unidad del edificio tiene
asignada una celda (campo "celda", ej. B12), la app NO crea ni modifica la planilla:
lee el mismo archivo tal cual está. La celda es la del nombre de la unidad y a su
derecha van Total a pagar, Monto deuda, Monto pagado y Tipo pago. Las unidades sin
celda no figuran en la planilla (no se puede generar su recibo).
"""

import os
import re
import shutil
import unicodedata
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

import config
from unidades import etiqueta_unidad, incluida_en_total, indice_por_id, ordenar_con_asociadas
from utils import parse_importe, sanitize_filename

ENCABEZADOS = ["Tipo unidad", "Total a pagar", "Monto deuda", "Monto pagado", "Tipo pago"]
FILA_TITULO, FILA_ENCABEZADO, FILA_INICIO = 2, 3, 4
COL_ETIQUETA, COL_TOTAL, COL_DEUDA, COL_PAGADO, COL_ESTADO, COL_ID = 2, 3, 4, 5, 6, 8

# Estado (tal como lo calcula la fórmula del Excel) -> texto de "Gastos de".
# None significa "usar el mes anterior".
TEXTO_POR_ESTADO = {
    "total": None,
    "a cta": "A CTA.",
    "parcial": "A CTA.",          # el Excel de una versión anterior decía "Parcial"
    "deuda": "DEUDA",
    "no pagado": "NO PAGADO",
}


# ---------------------------------------------------------------------------
# Rutas, etiquetas y orden
# ---------------------------------------------------------------------------

def ruta_pagos_edificio(nombre_edificio):
    return os.path.join(config.PAGOS_DIR, f"{sanitize_filename(nombre_edificio)}_pagos.xlsx")


def normalizar_celda(texto):
    """'b12' -> 'B12' ('' si está vacío). Lanza ValueError si no es una celda válida (letra/s + número)."""
    texto = str(texto or "").strip().replace(" ", "").upper()
    if not texto:
        return ""
    try:
        columna, fila = coordinate_from_string(texto)
        if fila < 1 or column_index_from_string(columna) > 16384:
            raise ValueError
    except Exception:
        raise ValueError(f"«{texto}» no es una celda válida. Escribí la columna y la fila, por ejemplo B12.") from None
    return f"{columna}{fila}"


def modo_celdas(unidades):
    """True si alguna unidad tiene celda asignada: la planilla del edificio se lee tal cual, sin tocarla."""
    return any(str(u.get("celda") or "").strip() for u in unidades)


# ---------------------------------------------------------------------------
# Escritura de la planilla
# ---------------------------------------------------------------------------

def renombrar_planilla(nombre_viejo, nombre_nuevo, actualizar_titulo=True):
    """
    Renombra el archivo de la planilla de un edificio y actualiza el título (B2), salvo
    que actualizar_titulo sea False (planilla propia en modo celdas: solo se renombra el
    archivo). Devuelve (ruta_vieja, ruta_nueva), o None si no había planilla. Lanza
    PermissionError si el archivo está abierto en Excel.
    """
    origen, destino = ruta_pagos_edificio(nombre_viejo), ruta_pagos_edificio(nombre_nuevo)
    if not os.path.isfile(origen):
        return None
    if os.path.exists(destino) and os.path.normcase(destino) != os.path.normcase(origen):
        raise ValueError("Ya existe una planilla de pagos con ese nombre.")
    os.replace(origen, destino)
    if not actualizar_titulo:
        return origen, destino
    try:
        wb = load_workbook(destino)
        try:
            wb.worksheets[0].cell(FILA_TITULO, 2).value = nombre_nuevo
            wb.save(destino)
        finally:
            wb.close()
    except Exception:
        os.replace(destino, origen)
        raise
    return origen, destino


def _borde(izq="thin", der="thin", arr="thin", aba="thin"):
    return Border(left=Side(style=izq), right=Side(style=der), top=Side(style=arr), bottom=Side(style=aba))


def _formula_estado(r):
    return f'=IF(E{r}=0,"No pagado",IF(E{r}=D{r},"Deuda",IF(E{r}>=C{r}+D{r},"Total","A cta.")))'


def _sin_espacios(texto):
    return str(texto or "").replace(" ", "").upper()


def _leer_existente(ruta):
    """Devuelve (orden_ids, etiquetas, valores, formulas_al_dia) de la planilla actual, o listas vacías."""
    wb = load_workbook(ruta, data_only=False)
    try:
        ws = wb.worksheets[0]
        orden, etiquetas, valores = [], [], {}
        formulas_al_dia = True
        if not any(ws.cell(r, COL_ID).value for r in range(FILA_INICIO, ws.max_row + 1)):
            # sin unidades: tiene que estar la fila de ejemplo con la fórmula (ver sincronizar_edificio)
            formulas_al_dia = _sin_espacios(ws.cell(FILA_INICIO, COL_ESTADO).value) == _sin_espacios(_formula_estado(FILA_INICIO))
        for r in range(FILA_INICIO, ws.max_row + 1):
            uid = ws.cell(r, COL_ID).value
            if not uid or uid in valores:
                continue
            uid = str(uid)
            orden.append(uid)
            etiquetas.append(ws.cell(r, COL_ETIQUETA).value)
            valores[uid] = (ws.cell(r, COL_TOTAL).value, ws.cell(r, COL_DEUDA).value, ws.cell(r, COL_PAGADO).value)
            formulas_al_dia = formulas_al_dia and _sin_espacios(ws.cell(r, COL_ESTADO).value) == _sin_espacios(_formula_estado(r))
        return orden, etiquetas, valores, formulas_al_dia
    finally:
        wb.close()


def _es_ajena(ruta):
    """
    True si el archivo es un Excel de la persona y no una planilla armada por esta app:
    la app siempre deja oculta la columna de ids (H) y este no, y además tiene datos.
    """
    wb = load_workbook(ruta, data_only=False)
    try:
        ws = wb.worksheets[0]
        if ws.column_dimensions["H"].hidden:
            return False
        return any(c.value not in (None, "") for fila in ws.iter_rows(min_row=FILA_INICIO) for c in fila)
    finally:
        wb.close()


def _respaldar(ruta):
    base, ext = os.path.splitext(ruta)
    copia = f"{base}_respaldo_{datetime.now():%Y%m%d_%H%M%S}{ext}"
    shutil.copy2(ruta, copia)
    return copia


class PlanillaAjena(Exception):
    """El archivo de la planilla tiene datos que no armó esta app (es el Excel propio de la persona)."""


def sincronizar_edificio(nombre_edificio, unidades, reemplazar_ajena=False):
    """
    Crea o actualiza la planilla del edificio para que tenga exactamente una
    fila por unidad, en orden, conservando los montos ya cargados.
    Devuelve True si escribió el archivo. Lanza PermissionError si el archivo
    está abierto en Excel. Si el archivo es un Excel propio (no lo armó la app) no lo
    toca y lanza PlanillaAjena, salvo reemplazar_ajena=True: entonces guarda antes una
    copia «..._respaldo_<fecha>.xlsx» y lo reemplaza.
    """
    ruta = ruta_pagos_edificio(nombre_edificio)
    os.makedirs(config.PAGOS_DIR, exist_ok=True)

    por_id = indice_por_id(unidades)
    # Las cocheras/bauleras incluidas en el total de su depto no tienen fila propia.
    deseadas = [u for u, _nivel in ordenar_con_asociadas(unidades) if not incluida_en_total(u)]
    ids = [u["id"] for u in deseadas]
    etiquetas = [etiqueta_unidad(u, por_id) for u in deseadas]

    valores = {}
    if os.path.isfile(ruta):
        if _es_ajena(ruta):
            if not reemplazar_ajena:
                raise PlanillaAjena(
                    "La planilla de pagos de este edificio tiene datos que no armó la app, así que no se modificó:\n"
                    f"{ruta}\n\n"
                    "Para usarla tal cual, asigná la celda de cada unidad (Administrar → «Asignar celdas»). "
                    "Si querés que la app arme la suya, renombrá o borrá ese archivo."
                )
            _respaldar(ruta)   # es un Excel de la persona (con su propio orden): no perderlo
        else:
            orden_actual, etiquetas_actuales, valores, formulas_al_dia = _leer_existente(ruta)
            if orden_actual == ids and etiquetas_actuales == etiquetas and formulas_al_dia:
                return False

    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja 1"
    fuente, fuente_b = Font(name="Arial"), Font(name="Arial", bold=True)
    centro = Alignment(horizontal="center", vertical="center")

    ws.merge_cells(start_row=FILA_TITULO, start_column=2, end_row=FILA_TITULO, end_column=6)
    for c in range(2, 7):
        celda = ws.cell(FILA_TITULO, c)
        celda.border = _borde("medium" if c == 2 else None, "medium" if c == 6 else None, "medium", "medium")
    ws.cell(FILA_TITULO, 2, nombre_edificio).font = fuente_b
    ws.cell(FILA_TITULO, 2).alignment = centro

    for i, titulo in enumerate(ENCABEZADOS):
        c = ws.cell(FILA_ENCABEZADO, 2 + i, titulo)
        c.font, c.alignment = fuente_b, centro
        c.border = _borde("medium" if i == 0 else "thin", "medium" if i == 4 else "thin", "medium", "medium")

    for i, u in enumerate(deseadas):
        r = FILA_INICIO + i
        ultima = i == len(deseadas) - 1
        anteriores = valores.get(u["id"], (None, None, None))
        contenido = [etiquetas[i], *anteriores, _formula_estado(r)]
        for j, valor in enumerate(contenido):
            c = ws.cell(r, 2 + j, valor)
            c.font, c.alignment = fuente, centro
            if 1 <= j <= 3:
                c.number_format = "#,##0.00"
            c.border = _borde("medium" if j == 0 else "thin", "medium" if j == 4 else "thin",
                              "medium" if i == 0 else "thin", "medium" if ultima else "thin")
        ws.cell(r, COL_ID, u["id"])

    if not deseadas:
        # Todavía sin unidades: deja igual la fórmula de "Tipo pago" en la primera fila, para poder copiarla.
        for j in range(5):
            c = ws.cell(FILA_INICIO, 2 + j, _formula_estado(FILA_INICIO) if j == 4 else None)
            c.font, c.alignment = fuente, centro
            if 1 <= j <= 3:
                c.number_format = "#,##0.00"
            c.border = _borde("medium" if j == 0 else "thin", "medium" if j == 4 else "thin", "thin", "medium")

    ws.column_dimensions["A"].width = 3.5
    ws.column_dimensions["B"].width = 22
    for letra in "CDEF":
        ws.column_dimensions[letra].width = 17
    ws.column_dimensions["H"].hidden = True
    ws.freeze_panes = "A4"
    wb.calculation.fullCalcOnLoad = True

    temporal = ruta + ".tmp"
    wb.save(temporal)
    try:
        os.replace(temporal, ruta)
    except Exception:
        if os.path.exists(temporal):
            os.remove(temporal)
        raise
    return True


# ---------------------------------------------------------------------------
# Lectura para generar recibos
# ---------------------------------------------------------------------------

def _calcular_estado(total, deuda, pagado):
    """Misma lógica que la fórmula de la columna 'Tipo pago' del Excel."""
    total, deuda, pagado = parse_importe(total), parse_importe(deuda), parse_importe(pagado)
    if pagado == 0:
        return "No pagado"
    if pagado == deuda:
        return "Deuda"
    if pagado >= total + deuda:
        return "Total"
    return "A cta."


def crear_planilla_en_blanco(nombre_edificio, unidades):
    """
    Rearma una planilla propia (modo celdas) que falta, en blanco. El nombre de cada unidad va en su
    celda y, a la derecha, Total a pagar, Monto deuda, Monto pagado y Tipo pago, con el mismo formato
    que las planillas automáticas. Si todas las unidades usan la misma columna, arriba de la primera
    fila va el título (nombre del edificio) y los encabezados, como en la planilla automática.
    Los montos anteriores no se recuperan. Solo están las unidades con celda asignada.
    """
    ruta = ruta_pagos_edificio(nombre_edificio)
    os.makedirs(config.PAGOS_DIR, exist_ok=True)
    por_id = indice_por_id(unidades)
    fuente, fuente_b = Font(name="Arial"), Font(name="Arial", bold=True)
    centro = Alignment(horizontal="center", vertical="center")

    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja 1"
    posiciones = []   # (unidad, fila, columna del nombre)
    for u in unidades:
        celda = str(u.get("celda") or "").strip()
        if celda:
            columna, fila = coordinate_from_string(celda)
            posiciones.append((u, fila, column_index_from_string(columna)))

    for u, fila, c0 in posiciones:
        total, deuda, pagado = (get_column_letter(c0 + i) for i in (1, 2, 3))
        formula = (f'=IF({pagado}{fila}=0,"No pagado",IF({pagado}{fila}={deuda}{fila},"Deuda",'
                   f'IF({pagado}{fila}>={total}{fila}+{deuda}{fila},"Total","A cta.")))')
        for i, valor in enumerate([etiqueta_unidad(u, por_id), None, None, None, formula]):
            c = ws.cell(fila, c0 + i, valor)
            c.font, c.alignment, c.border = fuente, centro, _borde()
            if 1 <= i <= 3:
                c.number_format = "#,##0.00"
        ws.column_dimensions[get_column_letter(c0)].width = 22
        for i in range(1, 5):
            ws.column_dimensions[get_column_letter(c0 + i)].width = 17

    columnas = {c0 for _u, _f, c0 in posiciones}
    primera = min((f for _u, f, _c in posiciones), default=0)
    if len(columnas) == 1 and primera >= 3:   # título y encabezados en las dos filas de arriba
        c0 = columnas.pop()
        fila_titulo, fila_encabezado = primera - 2, primera - 1
        ws.merge_cells(start_row=fila_titulo, start_column=c0, end_row=fila_titulo, end_column=c0 + 4)
        titulo = ws.cell(fila_titulo, c0, nombre_edificio)
        titulo.font, titulo.alignment = fuente_b, centro
        for i in range(5):
            ws.cell(fila_titulo, c0 + i).border = _borde()
            encabezado = ws.cell(fila_encabezado, c0 + i, ENCABEZADOS[i])
            encabezado.font, encabezado.alignment, encabezado.border = fuente_b, centro, _borde()

    temporal = ruta + ".tmp"
    wb.save(temporal)
    os.replace(temporal, ruta)
    return ruta


def planilla_abierta(nombre_edificio):
    """
    True si la planilla figura abierta en Excel o LibreOffice (existe su archivo de
    bloqueo). En ese caso lo que se escribió y no se guardó todavía no se ve desde la app.
    """
    ruta = ruta_pagos_edificio(nombre_edificio)
    carpeta, nombre = os.path.split(ruta)
    return os.path.exists(os.path.join(carpeta, "~$" + nombre)) or \
        os.path.exists(os.path.join(carpeta, ".~lock." + nombre + "#"))


def _fila_planilla(total, deuda, pagado, estado):
    if estado in (None, "") or str(estado).startswith(("=", "#")):
        estado = _calcular_estado(total, deuda, pagado)
    return {"total": parse_importe(total), "deuda": parse_importe(deuda),
            "pagado": parse_importe(pagado), "estado": str(estado).strip()}


def _leer_por_celdas(ws, unidades):
    """Lee cada unidad con celda asignada: la celda es el nombre y a su derecha total, deuda, pagado y tipo pago."""
    datos = {}
    for u in unidades:
        celda = str(u.get("celda") or "").strip()
        if not celda:
            continue
        columna, fila = coordinate_from_string(celda)
        c0 = column_index_from_string(columna)
        total, deuda, pagado, estado = (ws.cell(fila, c0 + i).value for i in range(1, 5))
        datos[str(u["id"])] = _fila_planilla(total, deuda, pagado, estado)
    return datos


def leer_planilla(nombre_edificio, unidades=None):
    """
    Devuelve {id_unidad: {"total", "deuda", "pagado", "estado"}} con lo que está
    GUARDADO en la planilla del edificio. Si el Excel no guardó el resultado de la
    fórmula "Tipo pago", el estado se calcula con la misma lógica. Lanza
    FileNotFoundError si la planilla no existe.
    Con 'unidades' (las del edificio) y alguna con celda asignada, lee por celdas (ver
    el modo "celdas" arriba); las unidades sin celda no figuran en el resultado.
    """
    ruta = ruta_pagos_edificio(nombre_edificio)
    if not os.path.isfile(ruta):
        raise FileNotFoundError(f"No se encontró la planilla de pagos:\n{ruta}")

    wb = load_workbook(ruta, data_only=True)
    try:
        ws = wb.worksheets[0]
        if unidades is not None and modo_celdas(unidades):
            return _leer_por_celdas(ws, unidades)
        datos = {}
        for r in range(FILA_INICIO, ws.max_row + 1):
            uid = ws.cell(r, COL_ID).value
            if not uid or str(uid) in datos:
                continue
            total, deuda, pagado = (ws.cell(r, c).value for c in (COL_TOTAL, COL_DEUDA, COL_PAGADO))
            datos[str(uid)] = _fila_planilla(total, deuda, pagado, ws.cell(r, COL_ESTADO).value)
        return datos
    finally:
        wb.close()


def leer_estados(nombre_edificio, unidades=None):
    """Devuelve {id_unidad: estado} (ver leer_planilla)."""
    return {uid: d["estado"] for uid, d in leer_planilla(nombre_edificio, unidades).items()}


def _norm(texto):
    t = unicodedata.normalize("NFKD", str(texto or ""))
    return re.sub(r"[^a-z0-9]+", " ", "".join(c for c in t if not unicodedata.combining(c)).lower()).strip()


def gastos_de_estado(estado, mes_anterior):
    """Texto de "Gastos de" para un estado de pago (None = no figura en la planilla: mes anterior)."""
    if estado is None:
        return mes_anterior
    texto = TEXTO_POR_ESTADO.get(_norm(estado), estado.upper())
    return mes_anterior if texto is None else texto


def gastos_de_unidad(estados, unidad, mes_anterior):
    """
    Devuelve (texto_gastos_de, estado). 'estado' es el "Tipo pago" de la
    planilla, o None si la unidad no figura (en ese caso: mes anterior).
    """
    estado = estados.get(str(unidad["id"]))
    return gastos_de_estado(estado, mes_anterior), estado


def estado_agregado(estados):
    """
    Estado de un grupo que paga todo junto (depto + cochera/baulera): "Total" solo si
    TODAS las filas dicen Total, "No pagado" si todas dicen No pagado, "Deuda" si todas
    dicen Deuda; cualquier otra combinación es "A cta.". None si no hay ninguna fila.
    """
    normalizados = [_norm(e) for e in estados if e]
    if not normalizados:
        return None
    for clave, texto in (("total", "Total"), ("no pagado", "No pagado"), ("deuda", "Deuda")):
        if all(e == clave for e in normalizados):
            return texto
    return "A cta."
